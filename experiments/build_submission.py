"""Build the Group 9 report and exactly four overview slides from recorded outputs."""

import argparse
import csv
import html
import json
import re
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "reports"
SETTINGS = ["centralized", "iid", "mild_non_iid", "strong_non_iid", "strong_e3"]
LABELS = ["Centralized", "FedAvg IID", "FedAvg mild", "FedAvg strong", "Strong E3R10"]


def rows_by_setting(path):
    with path.open() as handle:
        return {row["setting"]: row for row in csv.DictReader(handle)}


def table(headers, rows):
    return "\n".join(
        ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
        + ["| " + " | ".join(str(v) for v in row) + " |" for row in rows]
    )


def materialize():
    summary = rows_by_setting(ROOT / "outputs/tables/summary.csv")
    with (ROOT / "outputs/tables/per_seed.csv").open() as handle:
        seeds = {(row["setting"], int(row["seed"])): row for row in csv.DictReader(handle)}

    def mean(s, metric="accuracy"):
        return float(summary[s][metric + "_mean"])

    def fmt(s, metric="accuracy", scale=1, decimals=4):
        return f"{scale * mean(s, metric):.{decimals}f} +/- {scale * float(summary[s][metric + '_sd']):.{decimals}f}"

    cen, iid, mild, strong, e3 = [100 * mean(s) for s in SETTINGS]
    abstract = (
        "This project investigates how client label distributions affect federated image classification. "
        "A small convolutional neural network and a manually implemented FedAvg algorithm are evaluated "
        "on Fashion-MNIST with ten simulated clients. An exact quota design produces IID, mild and strong "
        "label skew while keeping client sample counts and global class totals fixed. Five configurations "
        "are evaluated over three seeds, with 1,620,000 training-image exposures per run. "
        f"Final test accuracy is {cen:.2f}% for centralized training, {iid:.2f}% for IID FedAvg, "
        f"{mild:.2f}% for mild skew and {strong:.2f}% for strong skew. "
        f"Under strong skew, three local epochs and ten rounds achieve {e3:.2f}% accuracy, "
        "close to one local epoch and thirty rounds at the same exposure budget. Error analysis identifies "
        "Shirt as the weakest class and examines traceable Shirt-to-T-shirt/top mistakes. "
        "Saved predictions, split indices and experiment manifests support an independent integrity audit. "
        "The findings demonstrate sensitivity to label skew within this controlled setting, while the local-epoch "
        "comparison suggests a potential aggregation trade-off. Conclusions remain limited to one dataset, "
        "one CNN, full participation and three seeds. The simulation supplies no formal privacy guarantee."
    )
    assert 150 <= len(abstract.split()) <= 200, len(abstract.split())
    manifest = json.loads((ROOT / "outputs/runs/centralized/seed_42/manifest.json").read_text())
    strong_metrics = [
        json.loads((ROOT / f"outputs/runs/strong_non_iid/seed_{s}/final_metrics.json").read_text())["final_test"]
        for s in [42, 43, 44]
    ]
    shirt = sum(m["per_class"][6]["recall"] for m in strong_metrics) / 3
    errors = strong_metrics[0]["confusion_matrix"][6][0]
    members = json.loads((REPORT / "members.json").read_text())
    contribution_rows = [[m["name"], m["student_id"], m["responsibility"]] for m in members["members"]]
    replacement = {
        "ABSTRACT": abstract,
        "RESULTS_TABLE": table(
            ["Setting", "Accuracy (%)", "Macro precision", "Macro-F1", "Test CE"],
            [
                [label, fmt(s, scale=100, decimals=2), fmt(s, "macro_precision"), fmt(s, "macro_f1"), fmt(s, "loss")]
                for s, label in zip(SETTINGS, LABELS)
            ],
        ),
        "ENVIRONMENT": (
            f"The recorded versions are PyTorch {manifest['versions']['torch']}, torchvision "
            f"{manifest['versions']['torchvision']}, NumPy {manifest['versions']['numpy']} and "
            f"scikit-learn {manifest['versions']['sklearn']}. The reproduction instructions target Python 3.12."
        ),
        "BASELINE_DISCUSSION": (
            f"Centralized accuracy is {cen:.2f}% compared with {iid:.2f}% for IID FedAvg, a gap of "
            f"{cen - iid:.2f} percentage points. Macro-F1 is {mean('centralized', 'macro_f1'):.4f} and "
            f"{mean('iid', 'macro_f1'):.4f}, respectively. Equal exposure therefore does not remove "
            "the difference between the two recorded optimization procedures."
        ),
        "SEVERITY_DISCUSSION": (
            f"Mild skew reduces mean accuracy by {iid - mild:.2f} percentage points relative to IID. "
            f"Strong skew reduces it by {iid - strong:.2f} points. The corresponding macro-F1 values "
            f"are {mean('mild_non_iid', 'macro_f1'):.4f} and {mean('strong_non_iid', 'macro_f1'):.4f}. "
            "The means decrease with severity across these three settings, but this finite comparison "
            "does not establish a universal monotonic relationship."
        ),
        "MILESTONE_TABLE": table(
            ["Setting", "Seeds reaching", "Mean point", "Effective epochs"],
            [
                [
                    label,
                    f"{summary[s]['r80_reached']} / 3",
                    f"{float(summary[s]['r80_round_mean']):.2f}" if summary[s]["r80_round_mean"] else "Not reached",
                    f"{float(summary[s]['r80_effective_epochs_mean']):.2f}"
                    if summary[s]["r80_effective_epochs_mean"]
                    else "Not reached",
                ]
                for s, label in zip(SETTINGS, LABELS)
            ],
        ),
        "ABLATION_DISCUSSION": (
            f"The E3R10 configuration reaches {e3:.2f}% test accuracy, compared with {strong:.2f}% "
            f"for E1R30. The difference is {e3 - strong:+.2f} percentage points. Their test cross "
            f"entropies are {mean('strong_e3', 'loss'):.4f} and {mean('strong_non_iid', 'loss'):.4f}, "
            f"and macro-F1 values are {mean('strong_e3', 'macro_f1'):.4f} and "
            f"{mean('strong_non_iid', 'macro_f1'):.4f}. The mean accuracy difference is much smaller "
            "than either setting's seed SD."
        ),
        "ABLATION_TABLE": table(
            ["Seed", "Strong E1R30 (%)", "Strong E3R10 (%)", "Difference (pp)"],
            [
                [
                    s,
                    f"{100 * float(seeds['strong_non_iid', s]['accuracy']):.2f}",
                    f"{100 * float(seeds['strong_e3', s]['accuracy']):.2f}",
                    f"{100 * (float(seeds['strong_e3', s]['accuracy']) - float(seeds['strong_non_iid', s]['accuracy'])):+.2f}",
                ]
                for s in [42, 43, 44]
            ],
        ),
        "ERROR_DISCUSSION": (
            f"For strong E1R30, Shirt has the lowest mean recall, {shirt:.4f} across the three seeds. "
            f"In seed 42, the largest off-diagonal confusion entry is Shirt predicted as T-shirt/top, "
            f"with {errors} of the 1,000 Shirt test images. This is an observed error pattern, "
            "not proof that label skew alone causes all Shirt mistakes."
        ),
        "CONCLUSION": (
            f"Within the fixed protocol, IID FedAvg trails the centralized reference by {cen - iid:.2f} "
            f"accuracy points. Strong label skew reduces FedAvg accuracy by {iid - strong:.2f} points "
            f"relative to IID. The E3R10 ablation gives {e3:.2f}% mean accuracy versus {strong:.2f}% "
            "for E1R30 with twenty fewer aggregation rounds at the same exposure budget. "
            "These findings answer the three research questions descriptively: optimization coordination "
            "matters, stronger local label concentration is associated with poorer global classification, "
            "and fewer aggregations can retain similar observed mean performance in this experiment."
        ),
        "CONTRIBUTIONS": (
            (
                "The following table records responsibilities for final integration and verification. "
                "The group confirmed these final responsibility assignments.\n\n"
                if members["confirmed"]
                else "Member identities are supplied by the group. Final responsibility assignments await group confirmation.\n\n"
            )
            + table(["Member", "Student ID", "Responsibility"], contribution_rows)
            + "\n\nGroup leader: "
            + members["leader"]
            + "."
        ),
        "PER_SEED_TABLE": table(
            ["Setting", "Seed 42", "Seed 43", "Seed 44", "Mean +/- SD"],
            [
                [label]
                + [f"{100 * float(seeds[s, z]['accuracy']):.2f}" for z in [42, 43, 44]]
                + [fmt(s, scale=100, decimals=2)]
                for s, label in zip(SETTINGS, LABELS)
            ],
        ),
        "BEST_VALIDATION_TABLE": table(
            ["Setting", "Mean test acc. (%)", "SD (%)", "Selected points 42/43/44"],
            [
                [
                    label,
                    f"{100 * float(summary[s]['best_validation_test_accuracy_mean']):.2f}",
                    f"{100 * float(summary[s]['best_validation_test_accuracy_sd']):.2f}",
                    "/".join(seeds[s, z]["best_validation_round"] for z in [42, 43, 44]),
                ]
                for s, label in zip(SETTINGS, LABELS)
            ],
        ),
    }
    source = (REPORT / "report_template_en.md").read_text()
    for key, value in replacement.items():
        source = source.replace("{{" + key + "}}", value)
    if re.search(r"\{\{.*?\}\}", source):
        raise ValueError("Unresolved report field")
    (REPORT / "9_27_Report.md").write_text(source)
    return source, summary, members, len(abstract.split())


def register_fonts():
    pairs = [
        (
            "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
            "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf",
        ),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf"),
    ]
    pair = next((p for p in pairs if all(Path(f).exists() for f in p)), None)
    if pair is None:
        raise RuntimeError("Install Times New Roman or DejaVu Serif before building the report")
    for name, file in zip(["Academic", "AcademicBold"], pair):
        pdfmetrics.registerFont(TTFont(name, file))
    pdfmetrics.registerFontFamily(
        "Academic", normal="Academic", bold="AcademicBold", italic="Academic", boldItalic="AcademicBold"
    )


def inline(text):
    escaped = html.escape(text)
    escaped = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"`(.*?)`", r'<font size="9">\1</font>', escaped)
    return re.sub(r"https?://[^\s<]+", lambda m: f'<link href="{m[0]}" color="#17386a">{m[0]}</link>', escaped)


def build_pdf(source, members):
    register_fonts()
    styles = {
        "body": ParagraphStyle(
            "Body", fontName="Academic", fontSize=10.8, leading=15, alignment=TA_JUSTIFY, spaceAfter=8
        ),
        "h1": ParagraphStyle("H1", fontName="AcademicBold", fontSize=16, leading=21, spaceAfter=13, keepWithNext=True),
        "h2": ParagraphStyle(
            "H2", fontName="AcademicBold", fontSize=12, leading=16, spaceBefore=7, spaceAfter=7, keepWithNext=True
        ),
        "cell": ParagraphStyle("Cell", fontName="Academic", fontSize=9, leading=12),
        "head": ParagraphStyle("Head", fontName="AcademicBold", fontSize=9, leading=12),
        "caption": ParagraphStyle(
            "Caption", fontName="Academic", fontSize=9.2, leading=12, spaceAfter=10, alignment=TA_CENTER
        ),
        "code": ParagraphStyle("Code", fontName="Courier", fontSize=8.2, leading=12, spaceAfter=8, leftIndent=8),
        "center": ParagraphStyle("Center", fontName="Academic", fontSize=12, leading=18, alignment=TA_CENTER),
        "cover": ParagraphStyle(
            "Cover", fontName="AcademicBold", fontSize=23, leading=30, alignment=TA_CENTER, spaceAfter=10
        ),
    }
    width = A4[0] - 120

    def para(text, style="body"):
        return Paragraph(inline(text), styles[style])

    def make_table(rows):
        n = len(rows[0])
        fractions = [1 / n] * n
        if n == 2:
            fractions = [0.85, 0.15]
        if n == 5:
            fractions = [0.23, 0.195, 0.195, 0.19, 0.19]
        if n == 6:
            fractions = [0.21, 0.10, 0.18, 0.18, 0.19, 0.14]
        if n == 3:
            fractions = [0.28, 0.18, 0.54]
        cells = [[para(c, "head" if r == 0 else "cell") for c in row] for r, row in enumerate(rows)]
        t = Table(cells, colWidths=[width * f for f in fractions], repeatRows=1, hAlign="CENTER")
        t.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f5")),
                    ("LINEABOVE", (0, 0), (-1, 0), 0.7, colors.black),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
                    ("LINEBELOW", (0, -1), (-1, -1), 0.7, colors.black),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        return [t, Spacer(1, 10)]

    flow = [
        Spacer(1, 35),
        para("VIETNAM ACADEMY OF SCIENCE AND TECHNOLOGY", "center"),
        para("UNIVERSITY OF SCIENCE AND TECHNOLOGY OF HANOI", "center"),
        Spacer(1, 12),
        para("Department of Information and Communication Technology", "center"),
        Spacer(1, 90),
        para("Deep Learning Final Project", "center"),
        Spacer(1, 12),
        Paragraph("Federated Image Classification<br/>under Non-IID Data", styles["cover"]),
        para("Group 9 - Topic 27", "center"),
        Spacer(1, 42),
    ]
    for m in members["members"]:
        flow.append(para(m["name"] + " - " + m["student_id"], "center"))
    flow += [Spacer(1, 55), para("Hanoi, October 2026", "center"), PageBreak()]
    flow += [para("Contents", "h1"), Spacer(1, 8)]
    contents = [
        ("Abstract; I. Introduction and Research Question", 1),
        ("II. Related Work", 2),
        ("III. Dataset and Data Preparation", 3),
        ("IV. Controlled Client Partitioning", 4),
        ("V. Methods", 5),
        ("V. Methods: FedAvg and implementation", 6),
        ("VI. Experimental Setup", 7),
        ("VI. Evaluation and reproducibility", 8),
        ("VII. Results and Discussion", 9),
        ("VII. Learning progress and stability", 10),
        ("VII. Local-epoch ablation", 11),
        ("VIII. Error and Qualitative Analysis", 12),
        ("IX. Conclusion and Limitations", 13),
        ("References", 14),
        ("Appendix A. Contributions; Appendix B. Per-seed results", 15),
        ("Appendix B. Reproduction and artifact index", 16),
    ]
    flow += make_table([["Section", "Page"]] + [[a, str(b)] for a, b in contents])
    flow += [
        Spacer(1, 20),
        para("Abbreviations", "h2"),
        para(
            "CNN: convolutional neural network. FL: federated learning. IID: independent and identically distributed. "
            "SGD: stochastic gradient descent. CE: cross entropy. TV: total variation. SD: sample standard deviation. "
            "E: local epochs. R: communication rounds. pp: percentage points."
        ),
        PageBreak(),
    ]
    chunks = source.split("<!-- page -->")
    if len(chunks) != 16:
        raise ValueError(f"Expected 13 main pages plus references and two appendix pages, got {len(chunks)}")
    for page_i, chunk in enumerate(chunks):
        lines = chunk.strip().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue
            if line.startswith("# "):
                flow.append(para(line[2:], "h1"))
            elif line.startswith("## "):
                flow.append(para(line[3:], "h2"))
            elif line.startswith("```"):
                code = []
                i += 1
                while i < len(lines) and not lines[i].startswith("```"):
                    code.append(html.escape(lines[i]).replace(" ", "&nbsp;"))
                    i += 1
                flow.append(Paragraph("<br/>".join(code), styles["code"]))
            elif line.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch("[-:]+", c) for c in cells):
                        rows.append(cells)
                    i += 1
                flow += make_table(rows)
                continue
            elif line.startswith("!["):
                m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
                im = Image(str(REPORT / m[2]))
                maxheight = 135 if page_i in (10, 11) else 185
                factor = min(width / im.imageWidth, maxheight / im.imageHeight)
                im.drawWidth = im.imageWidth * factor
                im.drawHeight = im.imageHeight * factor
                flow.append(KeepTogether([Spacer(1, 4), im, Spacer(1, 5), para(m[1], "caption")]))
            else:
                text = [line]
                while (
                    i + 1 < len(lines) and lines[i + 1].strip() and not lines[i + 1].startswith(("#", "|", "![", "```"))
                ):
                    i += 1
                    text.append(lines[i].strip())
                flow.append(para(" ".join(text), "caption" if line.startswith(("Table ", "Listing ")) else "body"))
            i += 1
        if page_i < len(chunks) - 1:
            flow.append(PageBreak())

    def footer(canvas, doc):
        if doc.page < 3:
            return
        canvas.saveState()
        canvas.setFont("Academic", 9)
        canvas.drawCentredString(A4[0] / 2, 30, str(doc.page - 2))
        canvas.restoreState()

    target = REPORT / "9_27_Report.pdf"
    doc = SimpleDocTemplate(
        str(target),
        pagesize=A4,
        leftMargin=60,
        rightMargin=60,
        topMargin=45,
        bottomMargin=48,
        title="Group 9 - Topic 27 - Federated Image Classification under Non-IID Data",
        author="Group 9",
        pageCompression=1,
    )
    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    if doc.page != 18:
        raise ValueError(f"Layout overflow: expected 18 physical pages, got {doc.page}")
    return target


def build_slides(summary):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    navy = RGBColor(25, 46, 85)

    def textbox(slide, x, y, w, h, text, size=23, bold=False):
        tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Arial"
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = navy
            p.space_after = Pt(18)
        return tf

    titles = ["Problem & Research Question", "Method & Experiments", "Key Results", "Conclusion / Demo"]
    notes = [
        "0:00-0:40. We study image classification when clients hold different label proportions. "
        "Ten simulated clients keep equal data quantities. Our questions concern the centralized gap, "
        "the effect of label skew and the trade-off between local epochs and aggregation rounds.",
        "0:40-1:25. All configurations use one CNN and the same train/validation/test split. "
        "We implement weighted FedAvg with a common round snapshot. The quota design gives lambda 0, "
        "0.5 and 0.9. Setup 1 is centralized versus IID. Setup 2 varies skew. Setup 3 compares strong "
        "E1R30 with E3R10 at equal exposure. Five configurations and three seeds give fifteen runs.",
        "1:25-2:15. Accuracy falls from centralized to IID and further under strong skew. "
        "The gap between IID and strong is 5.52 percentage points. E3R10 and E1R30 have similar "
        "observed means, but three seeds do not prove equivalence. The table uses final checkpoints "
        "and sample SD. No test-based checkpoint selection is used.",
        "2:15-2:55. The findings apply to this small CNN and synthetic label skew. "
        "Client drift is a plausible explanation, not directly measured. Shirt is the weakest class. "
        "For the demo, load pretrained artifacts before the exam, choose strong seed 42 and display "
        "one prediction. The demo is inference from a trained checkpoint. Stop by three minutes. "
        "The following twelve minutes are examiner questions for any group member.",
    ]
    slides = []
    for title, note in zip(titles, notes):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox(slide, 0.65, 0.4, 12, 0.65, title, 34, True)
        textbox(slide, 0.65, 6.98, 12, 0.3, "Group 9 - Topic 27", 12)
        slide.notes_slide.notes_text_frame.text = note
        slides.append(slide)
    textbox(slides[0], 0.7, 1.45, 11.9, 1.1, "Federated Image Classification\nunder Non-IID Data", 38, True)
    textbox(
        slides[0],
        0.7,
        3.1,
        11.8,
        2.9,
        "How does client label skew change global classification quality?\n"
        "10 clients, equal image counts, different label proportions.\n"
        "Compare centralized training, skew severity and local epochs.",
        26,
    )
    textbox(
        slides[1],
        0.7,
        1.45,
        12,
        1.1,
        "Fashion-MNIST: 54k train / 6k validation / 10k test\n"
        "Small CNN, 105,866 parameters. Weighted FedAvg in PyTorch.",
        24,
    )
    textbox(
        slides[1],
        0.7,
        3.0,
        11.8,
        3.5,
        "Setup 1: Centralized versus FedAvg IID\n"
        "Setup 2: IID, mild and strong label skew\n"
        "Setup 3: Strong E1R30 versus E3R10 at equal exposure\n"
        "5 configurations x 3 seeds. 1,620,000 image exposures per run.",
        25,
    )
    frame = slides[2].shapes.add_table(6, 3, Inches(0.75), Inches(1.45), Inches(7.3), Inches(3.85)).table
    values = [["Setting", "Accuracy (%)", "Macro-F1"]] + [
        [
            label,
            f"{100 * float(summary[s]['accuracy_mean']):.2f} +/- {100 * float(summary[s]['accuracy_sd']):.2f}",
            f"{float(summary[s]['macro_f1_mean']):.4f}",
        ]
        for s, label in zip(SETTINGS, LABELS)
    ]
    frame.columns[0].width = Inches(3)
    frame.columns[1].width = Inches(2.55)
    frame.columns[2].width = Inches(1.75)
    for r, row in enumerate(values):
        for c, text in enumerate(row):
            cell = frame.cell(r, c)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(238, 242, 247) if r == 0 else RGBColor(255, 255, 255)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Arial"
                p.font.size = Pt(20)
                p.font.bold = r == 0
                p.font.color.rgb = navy
    textbox(
        slides[2],
        8.55,
        1.75,
        3.95,
        3.8,
        "Strong skew: -5.52 pp versus IID\n\nE3R10: +0.04 pp versus E1R30 with 20 fewer rounds",
        26,
        True,
    )
    textbox(slides[2], 0.8, 5.7, 11.9, 0.6, "Final checkpoints. Mean +/- sample SD over three seeds.", 20)
    textbox(
        slides[3],
        0.7,
        1.4,
        7.0,
        4.8,
        "Label skew reduces global quality in this protocol.\n"
        "Fewer rounds retain similar observed mean accuracy at fixed exposure.\n"
        "Limits: one dataset/CNN, three seeds, simulated clients.\n"
        "Demo: strong seed 42, one test prediction from a trained checkpoint.",
        24,
    )
    slides[3].shapes.add_picture(
        str(REPORT / "figures/07_error_examples.png"), Inches(8.0), Inches(2.05), width=Inches(4.6)
    )
    textbox(slides[3], 8.05, 5.0, 4.55, 1.4, "Shirt recall: 0.2690\nShirt to T-shirt/top: 267 errors (seed 42)", 21)
    target = REPORT / "9_27_Overview.pptx"
    prs.save(target)
    assert len(prs.slides) == 4
    return target


def generate_report_figures():
    """Create legible print figures from recorded counts and per-class metrics."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    with plt.rc_context({"font.size": 13, "axes.labelsize": 13, "axes.titlesize": 14}):
        fig, axes = plt.subplots(1, 3, figsize=(10, 3.1), layout="constrained")
        for ax, severity, label in zip(axes, [0, 0.5, 0.9], ["IID", "Mild", "Strong"]):
            manifest = json.loads((ROOT / f"outputs/partitions/lambda_{severity}_seed_42/manifest.json").read_text())
            counts = np.asarray(manifest["train"]["counts"])
            im = ax.imshow(counts / 5400, vmin=0, vmax=0.5, cmap="Blues", aspect="equal")
            ax.set(title=label, xlabel="Class label", ylabel="Client")
            ax.set_xticks([0, 3, 6, 9])
            ax.set_yticks([0, 3, 6, 9])
        fig.colorbar(im, ax=axes, label="Local class proportion", shrink=0.85)
        fig.savefig(REPORT / "figures/01_label_skew_report.png", dpi=220)
        plt.close(fig)
        names = [
            "T-shirt/top",
            "Trouser",
            "Pullover",
            "Dress",
            "Coat",
            "Sandal",
            "Shirt",
            "Sneaker",
            "Bag",
            "Ankle boot",
        ]
        fig, ax = plt.subplots(figsize=(12, 3.0), layout="constrained")
        for setting, label in zip(SETTINGS, LABELS):
            metrics = [
                json.loads((ROOT / f"outputs/runs/{setting}/seed_{seed}/final_metrics.json").read_text())["final_test"]
                for seed in [42, 43, 44]
            ]
            recall = np.mean([[c["recall"] for c in m["per_class"]] for m in metrics], axis=0)
            ax.plot(names, recall, marker="o", markersize=4, label=label)
        ax.set(ylabel="Mean test recall", ylim=(0, 1.08))
        ax.grid(axis="y", alpha=0.2)
        ax.legend(ncol=5, loc="upper center", bbox_to_anchor=(0.5, 1.27), fontsize=12, frameon=False)
        ax.tick_params(axis="x", labelsize=12)
        fig.savefig(REPORT / "figures/07_recall_report.png", dpi=220)
        plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report-only", action="store_true")
    args = parser.parse_args()
    generate_report_figures()
    source, summary, members, words = materialize()
    pdf = build_pdf(source, members)
    if not args.report_only:
        build_slides(summary)
    print(f"Report: {pdf}. Abstract: {words} words. Overview: 4 slides.")


if __name__ == "__main__":
    main()
